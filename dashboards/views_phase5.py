"""
Phase 5 Views - Analytics, Certificates, Bulk Operations, Permissions
"""
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, FileResponse, HttpResponseForbidden, HttpResponse
from django.views.decorators.http import require_http_methods
from django.db.models import Avg, Count, Q
from django.utils import timezone
from django.core.paginator import Paginator
import csv
import io
from exams.models import Exam, ExamAttempt
from .models import (
    ExamAnalytics, StudentPerformance, AttemptHistory,
    GradingRubric, RubricCriteria, RubricScore, RubricGrade,
    ExamSchedule, Certificate, CertificateTemplate,
    BulkImportJob, BulkExportJob, Role, Permission, UserRole,
    QuestionBank, QuestionCategory, QuestionChoice
)
from .analytics import (
    update_exam_analytics, get_exam_statistics,
    get_student_performance_summary, get_class_analytics
)
from .certificates import (
    create_certificate, get_student_certificates,
    batch_generate_certificates, generate_certificate_pdf
)
from .bulk_operations import BulkImporter, BulkExporter
from .permissions import (
    user_has_permission, require_permission,
    get_user_permissions, initialize_default_roles,
    assign_role_to_user, remove_role_from_user
)
from .scheduler import (
    schedule_exam, get_schedule_info,
    reschedule_exam, cancel_exam_schedule
)
import json


# ==================== ANALYTICS VIEWS ====================

@login_required
def analytics_dashboard(request):
    """Main analytics dashboard"""
    if not user_has_permission(request.user, 'view_analytics'):
        return HttpResponseForbidden("You don't have permission to access analytics.")
    
    if request.user.role == 'TEACHER':
        exams = Exam.objects.filter(created_by=request.user)
    elif request.user.role == 'STUDENT':
        exams = Exam.objects.filter(school_class=request.user.student_class)
    else:
        exams = Exam.objects.all()
    
    # Annotate with attempt count
    exams = exams.annotate(
        attempt_count=Count('exam_attempts', filter=Q(exam_attempts__status='submitted'))
    ).order_by('-created_at')

    # Pagination
    paginator = Paginator(exams, 10) # Show 10 exams per page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'exams': page_obj,
        'total_exams': paginator.count,
    }
    
    return render(request, 'analytics/dashboard.html', context)


@login_required
def exam_analytics(request, exam_id):
    """Detailed analytics for a specific exam"""
    exam = get_object_or_404(Exam, id=exam_id)
    
    # Update analytics first
    analytics = update_exam_analytics(exam)
    stats = get_exam_statistics(exam)
    
    context = {
        'exam': exam,
        'analytics': analytics,
        'stats': stats,
        'question_stats': stats['question_stats'],
    }
    
    return render(request, 'analytics/exam_detail.html', context)


@login_required
def student_performance_detail(request, student_id):
    """View student performance summary"""
    from accounts.models import User
    student = get_object_or_404(User, id=student_id)
    
    # Check permission
    if request.user != student and request.user.role != 'TEACHER':
        return HttpResponseForbidden("You don't have permission to view this.")
    
    performance = get_student_performance_summary(student)
    
    context = {
        'student': student,
        'performance': performance,
    }
    
    return render(request, 'analytics/student_performance.html', context)


@login_required
@require_http_methods(["GET"])
def api_exam_statistics(request, exam_id):
    """API endpoint for exam statistics (JSON)"""
    exam = get_object_or_404(Exam, id=exam_id)
    
    update_exam_analytics(exam)
    analytics = ExamAnalytics.objects.get(exam=exam)
    
    return JsonResponse({
        'exam': exam.title,
        'total_attempts': analytics.total_attempts,
        'total_passed': analytics.total_passed,
        'average_score': analytics.average_score,
        'highest_score': analytics.highest_score,
        'lowest_score': analytics.lowest_score,
        'pass_rate': analytics.pass_rate,
        'average_time': analytics.average_time_taken,
    })


# ==================== GRADING RUBRIC VIEWS ====================

@login_required
def rubric_list(request):
    """List all grading rubrics"""
    if request.user.role == 'TEACHER':
        rubrics = GradingRubric.objects.filter(created_by=request.user)
    else:
        rubrics = GradingRubric.objects.all()
    
    context = {'rubrics': rubrics}
    return render(request, 'grading/rubric_list.html', context)


@login_required
def rubric_detail(request, rubric_id):
    """View rubric details with criteria"""
    rubric = get_object_or_404(GradingRubric, id=rubric_id)
    criteria = rubric.criteria.all()
    
    context = {
        'rubric': rubric,
        'criteria': criteria,
    }
    
    return render(request, 'grading/rubric_detail.html', context)


@login_required
def rubric_create(request):
    """Create new grading rubric"""
    if request.method == 'POST':
        name = request.POST.get('name')
        description = request.POST.get('description')
        total_points = int(request.POST.get('total_points', 100))
        
        rubric = GradingRubric.objects.create(
            name=name,
            description=description,
            total_points=total_points,
            created_by=request.user,
        )
        
        # Create criteria
        criteria_names = request.POST.getlist('criteria_name')
        criteria_max_points = request.POST.getlist('criteria_max_points')
        
        for i, name in enumerate(criteria_names):
            if name:
                RubricCriteria.objects.create(
                    rubric=rubric,
                    name=name,
                    description=request.POST.get(f'criteria_desc_{i}'),
                    max_points=int(criteria_max_points[i]),
                    order=i,
                )
        
        return redirect('dashboards:rubric_detail', rubric_id=rubric.id)
    
    return render(request, 'grading/rubric_create.html')


# ==================== EXAM SCHEDULING VIEWS ====================

@login_required
def schedule_exam_view(request, exam_id):
    """Schedule exam for automatic opening"""
    exam = get_object_or_404(Exam, id=exam_id)
    
    if not user_has_permission(request.user, 'schedule_exam'):
        return HttpResponseForbidden("You don't have permission to schedule exams.")
    
    if request.method == 'POST':
        scheduled_date_str = request.POST.get('scheduled_date')
        auto_open = request.POST.get('auto_open') == 'on'
        auto_close = request.POST.get('auto_close') == 'on'
        close_at_str = request.POST.get('close_at')
        notify_before = int(request.POST.get('notify_before_minutes', 15))
        
        from datetime import datetime
        scheduled_date = datetime.fromisoformat(scheduled_date_str)
        close_at = datetime.fromisoformat(close_at_str) if close_at_str else None
        
        schedule_exam(
            exam=exam,
            scheduled_date=scheduled_date,
            auto_open=auto_open,
            auto_close=auto_close,
            close_at=close_at,
            notify_before_minutes=notify_before,
        )
        
        return redirect('dashboards:exam_detail', exam_id=exam.id)
    
    schedule = get_schedule_info(exam)
    
    context = {
        'exam': exam,
        'schedule': schedule,
    }
    
    return render(request, 'exams/schedule.html', context)


# ==================== CERTIFICATE VIEWS ====================

@login_required
def certificate_list(request):
    """List certificates for current user"""
    # Auto-generate certificates for eligible student attempts
    if request.user.role == 'STUDENT':
        attempts = ExamAttempt.objects.filter(
            student=request.user,
            status='submitted',
            exam__is_published=True
        ).select_related('exam')
        
        for attempt in attempts:
            # Check if certificate already exists
            if not Certificate.objects.filter(student=request.user, exam=attempt.exam).exists():
                # Check if passed (using exam's passing_marks)
                if attempt.total_score >= attempt.exam.passing_marks:
                    create_certificate(request.user, attempt.exam, attempt)

    if request.user.role == 'STUDENT':
        certificates = Certificate.objects.filter(student=request.user)
    elif request.user.role == 'TEACHER':
        certificates = Certificate.objects.filter(exam__created_by=request.user)
    elif request.user.role == 'PARENT':
        certificates = Certificate.objects.filter(student__in=request.user.children.all())
    else:
        certificates = Certificate.objects.all()
    
    context = {'certificates': certificates}
    return render(request, 'certificates/list.html', context)


@login_required
def certificate_detail(request, certificate_id):
    """View certificate details"""
    cert = get_object_or_404(Certificate, id=certificate_id)
    
    # Check permission
    if (request.user != cert.student and 
        request.user.role != 'TEACHER' and
        not request.user.is_staff):
        return HttpResponseForbidden("You don't have permission to view this certificate.")
    
    context = {'certificate': cert}
    return render(request, 'certificates/detail.html', context)


@login_required
def certificate_download(request, certificate_id):
    """Download certificate PDF"""
    cert = get_object_or_404(Certificate, id=certificate_id)
    
    # Check permission
    if (request.user != cert.student and 
        request.user.role != 'TEACHER' and
        not request.user.is_staff):
        return HttpResponseForbidden("You don't have permission to download this certificate.")
    
    # Always regenerate the PDF to ensure the latest template is used
    generate_certificate_pdf(cert)
    cert.refresh_from_db()
    
    if cert.pdf_file:
        try:
            return FileResponse(
                cert.pdf_file.open('rb'),
                as_attachment=True,
                filename=f"{cert.certificate_number}.pdf"
            )
        except FileNotFoundError:
            # Should not happen since we just generated it, but just in case
            generate_certificate_pdf(cert)
            cert.refresh_from_db()
            if cert.pdf_file:
                return FileResponse(
                    cert.pdf_file.open('rb'),
                    as_attachment=True,
                    filename=f"{cert.certificate_number}.pdf"
                )

    from django.contrib import messages
    messages.error(request, "Could not generate certificate PDF. Please contact support.")
    return redirect('dashboards:certificate_detail', certificate_id=cert.id)


@login_required
def batch_issue_certificates(request, exam_id):
    """Issue certificates to all passing students"""
    exam = get_object_or_404(Exam, id=exam_id)
    
    if not user_has_permission(request.user, 'issue_certificate'):
        return HttpResponseForbidden("You don't have permission to issue certificates.")
    
    if request.method == 'POST':
        passing_percentage = float(request.POST.get('passing_percentage', 60))
        certificates = batch_generate_certificates(exam, passing_percentage)
        
        return JsonResponse({
            'success': True,
            'issued': len(certificates),
        })
    
    context = {'exam': exam}
    return render(request, 'certificates/batch_issue.html', context)



# ==================== BULK OPERATIONS VIEWS ====================

@login_required
def download_import_template(request, format_type):
    """Download template for bulk import"""
    template_type = request.GET.get('type', 'questions')  # 'questions' or 'students'

    if format_type not in ['csv', 'excel', 'word']:
        return HttpResponseForbidden("Invalid format")

    if template_type == 'students':
        headers = ['first_name', 'last_name', 'username', 'email', 'password']
        sample_rows = [
            ['Fatimah', 'Bello', 'fatimahbello', 'fatimah@school.com', 'changeme123'],
            ['Ahmad', 'Yusuf', 'ahmadyusuf', 'ahmad@school.com', 'changeme123'],
        ]
        filename_base = 'student_template'
        heading = 'Student Import Template'
        notes = []
    else:
        headers = [
            'text', 'type', 'marks',
            'choice_1', 'correct_1',
            'choice_2', 'correct_2',
            'choice_3', 'correct_3',
            'choice_4', 'correct_4',
        ]
        sample_rows = [
            ['What is 2+2?', 'objective', '1', '3', 'false', '4', 'true', '5', 'false', '6', 'false'],
            ['Explain photosynthesis.', 'subjective', '5', '', '', '', '', '', '', '', ''],
        ]
        filename_base = 'question_template'
        heading = 'Question Import Template'
        notes = [
            'type values: objective | subjective',
            'correct_1 to correct_4: true or false (case-insensitive)',
            'Objective questions must have at least one correct answer marked true',
            'For subjective: leave all choice and correct columns empty',
        ]

    if format_type == 'csv':
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}.csv"'
        writer = csv.writer(response)
        writer.writerow(headers)
        for row in sample_rows:
            writer.writerow(row)
        if notes:
            writer.writerow([])
            for note in notes:
                writer.writerow([f'# {note}'])
        return response

    elif format_type == 'excel':
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = heading
        # Headers row — bold
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='D9E1F2')
        # Sample data rows
        for row in sample_rows:
            ws.append(row)
        # Notes section below data
        if notes:
            ws.append([])
            ws.append(['NOTES (read only — do not include in import):'])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color='FF0000')
            for note in notes:
                ws.append([f'  • {note}'])
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}.xlsx"'
        wb.save(response)
        return response

    elif format_type == 'word':
        import docx
        from docx.shared import Pt, RGBColor
        doc = docx.Document()
        doc.add_heading(heading, 0)
        if notes:
            doc.add_heading('Notes', level=2)
            for note in notes:
                doc.add_paragraph(f'• {note}', style='List Bullet')
        doc.add_heading('Template Table', level=2)
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].bold = True
        for row_data in sample_rows:
            row_cells = table.add_row().cells
            for i, val in enumerate(row_data):
                row_cells[i].text = str(val)
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}.docx"'
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        response.write(buffer.getvalue())
        return response

    return HttpResponseForbidden("Unknown error")

@login_required
def bulk_import_view(request):
    """Handle bulk imports"""
    if not user_has_permission(request.user, 'bulk_import'):
        return HttpResponseForbidden("You don't have permission to perform bulk imports.")
    
    if request.method == 'POST':
        import_type = request.POST.get('import_type')
        csv_file = request.FILES.get('csv_file')

        if not csv_file:
            from django.contrib import messages
            messages.error(request, 'No file uploaded.')
            return redirect('dashboards:bulk_import')

        if import_type == 'students':
            school_class_id = request.POST.get('school_class_id')
            from exams.models import SchoolClass
            school_class = SchoolClass.objects.get(id=school_class_id)
            job = BulkImporter.import_students(csv_file, school_class, request.user)

        elif import_type == 'questions':
            from exams.models import Exam as ExamModel
            from django.utils.dateparse import parse_datetime
            from django.utils.timezone import make_aware
            import datetime as dt
            exam_id = request.POST.get('exam_id')
            exam = ExamModel.objects.get(id=exam_id)
            start_time = request.POST.get('start_time')
            end_time = request.POST.get('end_time')
            publish = request.POST.get('publish_exam') == 'on'
            if start_time:
                parsed = parse_datetime(start_time)
                exam.start_time = make_aware(parsed) if parsed and parsed.tzinfo is None else parsed
            if end_time:
                parsed = parse_datetime(end_time)
                exam.end_time = make_aware(parsed) if parsed and parsed.tzinfo is None else parsed
            if publish:
                exam.is_published = True
            if start_time or end_time or publish:
                exam.save()
            job = BulkImporter.import_questions(csv_file, request.user, exam)
        else:
            from django.contrib import messages
            messages.error(request, 'Unknown import type.')
            return redirect('dashboards:bulk_import')

        return redirect('dashboards:bulk_import_job_detail', job_id=job.id)

    from exams.models import Exam as ExamModel, SchoolClass
    context = {
        'classes': SchoolClass.objects.all(),
        'exams': ExamModel.objects.select_related('school_class').order_by('-created_at'),
    }
    
    return render(request, 'bulk/import.html', context)


@login_required
def bulk_import_job_detail(request, job_id):
    """View bulk import job status"""
    job = get_object_or_404(BulkImportJob, id=job_id)
    
    context = {'job': job}
    return render(request, 'bulk/import_detail.html', context)


@login_required
def bulk_export_view(request):
    """Handle bulk exports"""
    if not user_has_permission(request.user, 'bulk_export'):
        return HttpResponseForbidden("You don't have permission to perform bulk exports.")
    
    if request.method == 'POST':
        export_type = request.POST.get('export_type')
        exam_id = request.POST.get('exam_id')
        file_format = request.POST.get('file_format', 'csv')
        
        exam = get_object_or_404(Exam, id=exam_id) if exam_id else None
        
        if export_type == 'results':
            job = BulkExporter.export_exam_results(exam, file_format)
        
        job.exported_by = request.user
        job.save(update_fields=['exported_by'])
        
        return redirect('dashboards:bulk_export_job_detail', job_id=job.id)
    
    context = {'exams': Exam.objects.all()}
    return render(request, 'bulk/export.html', context)


@login_required
def bulk_export_job_detail(request, job_id):
    """View bulk export job and download file"""
    job = get_object_or_404(BulkExportJob, id=job_id)
    
    if request.GET.get('download') == '1' and job.export_file:
        return FileResponse(
            job.export_file.open('rb'),
            as_attachment=True,
            filename=job.export_file.name.split('/')[-1],
        )
    
    context = {'job': job}
    return render(request, 'bulk/export_detail.html', context)


# ==================== PERMISSION & ROLE VIEWS ====================

@login_required
def role_management(request):
    """Manage roles and permissions"""
    if not request.user.is_staff:
        return HttpResponseForbidden("You don't have permission to manage roles.")
    
    roles = Role.objects.all()
    context = {'roles': roles}
    return render(request, 'admin/role_management.html', context)


@login_required
def assign_role_view(request, user_id):
    """Assign roles to users"""
    from accounts.models import User
    user = get_object_or_404(User, id=user_id)
    
    if not request.user.is_staff:
        return HttpResponseForbidden("You don't have permission to assign roles.")
    
    if request.method == 'POST':
        role_id = request.POST.get('role_id')
        expires_at_str = request.POST.get('expires_at')
        
        role = Role.objects.get(id=role_id)
        expires_at = None
        if expires_at_str:
            from datetime import datetime
            expires_at = datetime.fromisoformat(expires_at_str)
        
        assign_role_to_user(user, role, request.user, expires_at)
        
        return redirect('admin:auth_user_change', user.id)
    
    roles = Role.objects.all()
    context = {
        'user': user,
        'roles': roles,
        'current_roles': user.assigned_roles.all(),
    }
    
    return render(request, 'admin/assign_role.html', context)


@login_required
def permission_list(request):
    """List all permissions"""
    if not request.user.is_staff:
        return HttpResponseForbidden("You don't have permission to view permissions.")
    
    permissions = Permission.objects.all()
    context = {'permissions': permissions}
    return render(request, 'admin/permissions.html', context)


# ==================== QUESTION BANK VIEWS ====================

@login_required
def question_bank_list(request):
    """List all questions in the question bank"""
    questions = QuestionBank.objects.all()
    
    # Filter by category if provided
    category_id = request.GET.get('category')
    if category_id:
        questions = questions.filter(category_id=category_id)
    
    # Filter by difficulty
    difficulty = request.GET.get('difficulty')
    if difficulty:
        questions = questions.filter(difficulty=difficulty)
    
    # Search
    search = request.GET.get('search')
    if search:
        questions = questions.filter(
            Q(text__icontains=search) |
            Q(explanation__icontains=search)
        )
    
    context = {
        'questions': questions,
        'categories': QuestionCategory.objects.all(),
        'selected_category': category_id,
    }
    
    return render(request, 'question_bank/list.html', context)


@login_required
def question_bank_detail(request, question_id):
    """View question details"""
    question = get_object_or_404(QuestionBank, id=question_id)
    
    context = {'question': question}
    return render(request, 'question_bank/detail.html', context)


@login_required
def question_bank_create(request):
    """Create new question"""
    if not user_has_permission(request.user, 'create_question'):
        return HttpResponseForbidden("You don't have permission to create questions.")
    
    if request.method == 'POST':
        # Handle form submission
        text = request.POST.get('text')
        q_type = request.POST.get('type')
        marks = int(request.POST.get('marks', 1))
        difficulty = request.POST.get('difficulty', 'medium')
        category_id = request.POST.get('category')
        explanation = request.POST.get('explanation')
        
        category = QuestionCategory.objects.get(id=category_id) if category_id else None
        
        question = QuestionBank.objects.create(
            text=text,
            question_type=q_type,
            marks=marks,
            difficulty=difficulty,
            category=category,
            explanation=explanation,
            created_by=request.user,
            is_published=True,
        )
        
        # Handle choices for objective questions
        if q_type == 'objective':
            correct_choice_idx = request.POST.get('correct_choice')
            
            # Find all choice text fields
            for key, value in request.POST.items():
                if key.startswith('choice_text_') and value.strip():
                    try:
                        idx = key.split('_')[-1]
                        is_correct = (idx == correct_choice_idx)
                        
                        QuestionChoice.objects.create(
                            question=question,
                            text=value,
                            is_correct=is_correct,
                            order=int(idx)
                        )
                    except (ValueError, IndexError):
                        continue
        
        return redirect('dashboards:question_bank_detail', question_id=question.id)
    
    context = {'categories': QuestionCategory.objects.all()}
    return render(request, 'question_bank/create.html', context)


@login_required
def question_bank_edit(request, question_id):
    """Edit existing question"""
    question = get_object_or_404(QuestionBank, id=question_id)
    
    if not user_has_permission(request.user, 'edit_question'):
        return HttpResponseForbidden("You don't have permission to edit questions.")
    
    if request.method == 'POST':
        # Handle form submission
        question.text = request.POST.get('text', question.text)
        question.question_type = request.POST.get('type', question.question_type)
        question.marks = int(request.POST.get('marks', question.marks))
        question.difficulty = request.POST.get('difficulty', question.difficulty)
        question.explanation = request.POST.get('explanation', question.explanation)
        
        category_id = request.POST.get('category')
        if category_id:
            question.category_id = category_id
        
        question.save()

        # Handle choices for objective questions
        if question.question_type == 'objective':
            # Remove existing choices
            question.choices.all().delete()
            
            correct_choice_idx = request.POST.get('correct_choice')
            
            # Find all choice text fields
            for key, value in request.POST.items():
                if key.startswith('choice_text_') and value.strip():
                    try:
                        idx = key.split('_')[-1]
                        is_correct = (idx == correct_choice_idx)
                        
                        QuestionChoice.objects.create(
                            question=question,
                            text=value,
                            is_correct=is_correct,
                            order=int(idx)
                        )
                    except (ValueError, IndexError):
                        continue
        
        return redirect('dashboards:question_bank_detail', question_id=question.id)
    
    context = {
        'question': question,
        'categories': QuestionCategory.objects.all(),
    }
    return render(request, 'question_bank/edit.html', context)


@login_required
def rubric_edit(request, rubric_id):
    """Edit existing rubric"""
    rubric = get_object_or_404(GradingRubric, id=rubric_id)
    
    if not user_has_permission(request.user, 'edit_rubric'):
        return HttpResponseForbidden("You don't have permission to edit rubrics.")
    
    if request.method == 'POST':
        # Handle form submission
        rubric.name = request.POST.get('name', rubric.name)
        rubric.description = request.POST.get('description', rubric.description)
        rubric.total_points = int(request.POST.get('total_points', rubric.total_points))
        rubric.is_published = request.POST.get('is_published') == 'on'
        rubric.save()
        
        return redirect('dashboards:rubric_detail', rubric_id=rubric.id)
    
    context = {
        'rubric': rubric,
        'criteria': rubric.criteria.all(),
    }
    return render(request, 'grading/rubric_edit.html', context)
